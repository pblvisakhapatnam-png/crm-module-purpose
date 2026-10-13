"""Build the two user-test Feedback Logs from the testers' sheets (the restart of 12 October 2026).

Each tester records Pass or Fail for every step of their sheet (owner's decision, 9 October 2026), so each pack's
Feedback Log has a tab for each tester's sheet and one for their own data, listing every step in order, beside the
Findings tab - one row per problem - and a Summary. A step is a heading `### <n>. <title>` in a sheet; the tabs are
built from those headings, so a sheet whose steps change needs its log built again:

    python scripts/build_uat_feedback_logs.py          # writes docs/uat/... and docs/uat-hiring/..._FEEDBACK_LOG.xlsx

Needs `openpyxl` (pip), only to write. Reading the steps is standard library, so the tests
(scripts/tests/test_uat_tester_sheets.py) check a committed log against its sheets without it.
"""

from __future__ import annotations

import re
import sys
from dataclasses import dataclass
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[1]

STEP_HEADING = re.compile(r"^### (\d+)\. (.+?)\s*$", re.MULTILINE)
RESULTS = ("Pass", "Fail", "Could not try", "Not done")
AGAIN = ("yes, same thing", "yes, but different", "no", "did not try again")
IMPACT = ("stopped me", "slowed me down", "annoying", "just noting it")
LOGINS = ("uat-sales-01", "uat-manager-01", "uat-operations-01", "uat-finance-01", "uat-fleet-01", "uat-owner-01", "uat-sales-02")


@dataclass(frozen=True)
class Pack:
    folder: Path
    title: str
    log: str
    tabs: tuple[tuple[str, str], ...]  # (tab name, sheet file)

    @property
    def log_path(self) -> Path:
        return self.folder / self.log


PACKS = (
    Pack(
        REPO_ROOT / "docs" / "uat",
        "PBL CRM test - Feedback Log",
        "PBL_CRM_UAT_FEEDBACK_LOG.xlsx",
        (
            ("Tester 1 - Sales", "PBL_CRM_UAT_TESTER_1_SALES.md"),
            ("Tester 2 - Sales Manager", "PBL_CRM_UAT_TESTER_2_SALES_MANAGER.md"),
            ("Tester 3 - Operations", "PBL_CRM_UAT_TESTER_3_OPERATIONS.md"),
            ("Tester 4 - Accounts", "PBL_CRM_UAT_TESTER_4_ACCOUNTS.md"),
            ("Own data", "PBL_CRM_UAT_YOUR_OWN_DATA.md"),
        ),
    ),
    Pack(
        REPO_ROOT / "docs" / "uat-hiring",
        "PBL Hiring user test - Feedback Log",
        "PBL_HIRING_UAT_FEEDBACK_LOG.xlsx",
        (
            ("Tester 1 - Sales", "PBL_HIRING_UAT_TESTER_1_SALES.md"),
            ("Tester 2 - Fleet Admin", "PBL_HIRING_UAT_TESTER_2_FLEET_ADMIN.md"),
            ("Tester 3 - Operations", "PBL_HIRING_UAT_TESTER_3_OPERATIONS.md"),
            ("Tester 4 - Accounts", "PBL_HIRING_UAT_TESTER_4_ACCOUNTS.md"),
            ("Own data", "PBL_HIRING_UAT_YOUR_OWN_DATA.md"),
        ),
    ),
)


def steps(sheet: Path) -> list[tuple[int, str]]:
    """The sheet's steps, in order, as (number, title) - the plain words of each `### <n>. <title>` heading."""
    return [(int(number), title.replace("**", "").replace("`", "")) for number, title in STEP_HEADING.findall(sheet.read_text(encoding="utf-8"))]


HOW_TO_USE = [
    "How to fill in this log",
    "",
    "1. Open the tab for your tester number - or 'Own data' for the last part. Every step of your sheet is listed.",
    "2. After each step, choose its Result: Pass, Fail, Could not try (an earlier step failed, or another tester has",
    "   not got there yet - say which in Notes), or Not done.",
    "3. For a Fail, write a row in Findings for the problem, and put that row's number (#) beside the step.",
    "4. Findings: one row per problem, at the time - not saved up for the end. Four columns matter most: the login you",
    "   used, roughly what time it was, the record's code (CO-, WO-, RR-, AL-, EQ-) or name, and whether it happened again.",
    "5. If something looks badly wrong - data lost, a record you should not see, something your job should not allow -",
    "   tell the person running the test at once, and leave the screen as it is.",
    "",
    "Doing every sheet yourself? Fill in every tester's tab, in the order of the one-person route.",
    "Summary counts the results and the findings as you go.",
]

FINDINGS_COLUMNS = [
    ("#", 5), ("Your name", 16), ("Login used", 18), ("Sheet and step", 18), ("Date", 12), ("Time (roughly)", 10),
    ("Which screen", 24), ("Which record", 16), ("What you were trying to do", 32), ("What you did, step by step", 40),
    ("What you expected", 32), ("What happened instead", 40), ("Did it happen again?", 18),
    ("If different, what was different", 28), ("How much did it get in your way?", 18), ("Screenshot attached?", 12),
    ("For us - status", 14), ("For us - notes", 28),
]
FINDINGS_ROWS = 400


def build(pack: Pack) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    header_fill = PatternFill("solid", fgColor="E6F2F0")
    bold = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")

    def choices(values: tuple[str, ...]) -> str:
        return '"' + ",".join(values) + '"'

    workbook = Workbook()
    guide = workbook.active
    guide.title = "How to use"
    guide.column_dimensions["A"].width = 118
    for row, line in enumerate([pack.title, ""] + HOW_TO_USE, start=1):
        guide.cell(row=row, column=1, value=line)
    guide["A1"].font = Font(bold=True, size=14)

    step_tabs = []
    for tab_name, sheet_name in pack.tabs:
        sheet_steps = steps(pack.folder / sheet_name)
        tab = workbook.create_sheet(tab_name)
        step_tabs.append((tab_name, len(sheet_steps)))
        for column, (heading, width) in enumerate(
            (("Step", 6), ("What the step is", 52), ("Result", 14), ("Findings row(s)", 14), ("Notes", 60)), start=1
        ):
            cell = tab.cell(row=1, column=column, value=heading)
            cell.font, cell.fill = bold, header_fill
            tab.column_dimensions[cell.column_letter].width = width
        for row, (number, title) in enumerate(sheet_steps, start=2):
            tab.cell(row=row, column=1, value=number)
            tab.cell(row=row, column=2, value=title).alignment = wrap
            tab.cell(row=row, column=5).alignment = wrap
        last = len(sheet_steps) + 1
        result = DataValidation(type="list", formula1=choices(RESULTS), allow_blank=True)
        tab.add_data_validation(result)
        result.add(f"C2:C{last}")
        tab.freeze_panes = "A2"
        tab.cell(row=last + 2, column=2, value=f"From {sheet_name}").font = Font(italic=True, color="5B6776")

    findings = workbook.create_sheet("Findings")
    for column, (heading, width) in enumerate(FINDINGS_COLUMNS, start=1):
        cell = findings.cell(row=1, column=column, value=heading)
        cell.font, cell.fill, cell.alignment = bold, header_fill, wrap
        findings.column_dimensions[cell.column_letter].width = width
    for row in range(2, FINDINGS_ROWS + 1):
        findings.cell(row=row, column=1, value=row - 1)
    for column, values in ((3, LOGINS), (13, AGAIN), (15, IMPACT), (16, ("yes", "no"))):
        validation = DataValidation(type="list", formula1=choices(values), allow_blank=True)
        findings.add_data_validation(validation)
        letter = findings.cell(row=1, column=column).column_letter
        validation.add(f"{letter}2:{letter}{FINDINGS_ROWS}")
    findings.freeze_panes = "A2"

    summary = workbook.create_sheet("Summary")
    summary.column_dimensions["A"].width = 30
    for column in "BCDEF":
        summary.column_dimensions[column].width = 14
    summary["A1"] = "How the test is going"
    summary["A1"].font = Font(bold=True, size=14)
    summary["A2"] = "The counts update themselves as results and findings are filled in."
    row = 4
    for column, heading in enumerate(("Tab", *RESULTS, "Steps"), start=1):
        summary.cell(row=row, column=column, value=heading).font = bold
    for tab_name, count in step_tabs:
        row += 1
        summary.cell(row=row, column=1, value=tab_name)
        for column, result in enumerate(RESULTS, start=2):
            summary.cell(row=row, column=column, value=f"=COUNTIF('{tab_name}'!$C$2:$C${count + 1},\"{result}\")")
        summary.cell(row=row, column=len(RESULTS) + 2, value=count)
    row += 2
    summary.cell(row=row, column=1, value="Findings, by how much they got in the way").font = bold
    for impact in IMPACT:
        row += 1
        summary.cell(row=row, column=1, value=impact)
        summary.cell(row=row, column=2, value=f"=COUNTIF(Findings!$O$2:$O${FINDINGS_ROWS},$A{row})")
    row += 2
    summary.cell(row=row, column=1, value="Findings, by login").font = bold
    for login in LOGINS:
        row += 1
        summary.cell(row=row, column=1, value=login)
        summary.cell(row=row, column=2, value=f"=COUNTIF(Findings!$C$2:$C${FINDINGS_ROWS},$A{row})")

    workbook.save(pack.log_path)


def main() -> int:
    for pack in PACKS:
        build(pack)
        print("Wrote " + str(pack.log_path.relative_to(REPO_ROOT)))
    return 0


if __name__ == "__main__":
    sys.exit(main())
